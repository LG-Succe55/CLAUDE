#!/usr/bin/env python3
"""
Chrysler at CES — Concept Portfolio deck generator.

Premium automotive executive scaffold (16:9). Six concepts presented as a
single portfolio and a single keynote story, then each concept expanded into
intro / exterior / exterior innovations / interior / interior innovations /
advanced color & materials (ACM). White backgrounds, large imagery, minimal
text, subtle Chrysler-inspired accents, speaker-note placeholders throughout.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ----------------------------------------------------------------------------
# Design system
# ----------------------------------------------------------------------------
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

INK      = RGBColor(0x1A, 0x1A, 0x1C)
GRAPHITE = RGBColor(0x4A, 0x4E, 0x57)
SLATE    = RGBColor(0x86, 0x8B, 0x94)
PLATINUM = RGBColor(0xC8, 0xCC, 0xD2)
MIST     = RGBColor(0xEF, 0xF1, 0xF3)
CLOUD    = RGBColor(0xF7, 0xF8, 0xF9)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
CHRYSLER = RGBColor(0x1B, 0x2A, 0x4A)
SILVER   = RGBColor(0xA9, 0xAE, 0xB5)

FONT_DISPLAY = "Arial"
FONT_BODY    = "Arial"

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
BLANK = prs.slide_layouts[6]

_PAGE = [0]
def page():
    _PAGE[0] += 1
    return _PAGE[0]


# ----------------------------------------------------------------------------
# Helpers
# ----------------------------------------------------------------------------
def add_slide():
    return prs.slides.add_slide(BLANK)


def set_bg(slide, color=WHITE):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color


def _set_font(run, size, color, bold, name, spacing, italic=False):
    f = run.font
    f.size = Pt(size); f.bold = bold; f.italic = italic
    f.name = name; f.color.rgb = color
    if spacing is not None:
        run._r.get_or_add_rPr().set("spc", str(int(spacing * 100)))


def textbox(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ln.get("align", align)
        if ln.get("space_after") is not None:
            p.space_after = Pt(ln["space_after"])
        if ln.get("space_before") is not None:
            p.space_before = Pt(ln["space_before"])
        if ln.get("line_spacing") is not None:
            p.line_spacing = ln["line_spacing"]
        r = p.add_run()
        r.text = ln["text"]
        _set_font(r, ln.get("size", 18), ln.get("color", INK),
                  ln.get("bold", False), ln.get("name", FONT_BODY),
                  ln.get("spacing"), ln.get("italic", False))
    return tb


def rect(slide, x, y, w, h, fill=None, line=None, line_w=None,
         shape=MSO_SHAPE.RECTANGLE):
    sp = slide.shapes.add_shape(shape, x, y, w, h)
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = line_w or Pt(0.75)
    return sp


def hairline(slide, x, y, w, color=PLATINUM, weight=Pt(1.0)):
    ln = slide.shapes.add_connector(2, x, y, x + w, y)
    ln.line.color.rgb = color
    ln.line.width = weight
    return ln


def kicker(slide, x, y, text, color=CHRYSLER, w=Inches(9)):
    return textbox(slide, x, y, w, Inches(0.3),
                   [{"text": text.upper(), "size": 11, "color": color,
                     "bold": True, "name": FONT_DISPLAY, "spacing": 3.0}])


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = "SPEAKER NOTES (placeholder) — " + text


def footer(slide, number):
    textbox(slide, Inches(0.7), Inches(7.04), Inches(8), Inches(0.3),
            [{"text": "CHRYSLER  ·  CES CONCEPT PORTFOLIO", "size": 8,
              "color": SLATE, "spacing": 1.5}])
    textbox(slide, Inches(11.6), Inches(7.04), Inches(1.0), Inches(0.3),
            [{"text": f"{number:02d}", "size": 8, "color": SLATE, "spacing": 1.5}],
            align=PP_ALIGN.RIGHT)
    textbox(slide, Inches(11.4), Inches(0.34), Inches(1.2), Inches(0.3),
            [{"text": "CONFIDENTIAL", "size": 7.5, "color": PLATINUM,
              "spacing": 2.0}], align=PP_ALIGN.RIGHT)


def image_frame(slide, x, y, w, h, label="IMAGE PLACEHOLDER", fill=MIST,
                note="Replace · 16:9 · 300dpi", icon=True):
    rect(slide, x, y, w, h, fill=fill)
    tick = Inches(0.26); tw = Pt(1.25)
    for cxp, cyp, pos in [(x, y, "tl"), (x + w, y, "tr"),
                          (x, y + h, "bl"), (x + w, y + h, "br")]:
        dx = tick if "l" in pos else -tick
        dy = tick if "t" in pos else -tick
        a = slide.shapes.add_connector(2, cxp, cyp, cxp + dx, cyp)
        a.line.color.rgb = SILVER; a.line.width = tw
        b = slide.shapes.add_connector(2, cxp, cyp, cxp, cyp + dy)
        b.line.color.rgb = SILVER; b.line.width = tw
    if icon and h > Inches(1.3):
        ic = Inches(0.56)
        icx = x + (w - ic) / 2
        icy = y + (h - ic) / 2 - Inches(0.24)
        rect(slide, icx, icy, ic, ic, fill=None, line=SILVER, line_w=Pt(1.5),
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        rect(slide, icx + ic * 0.22, icy + ic * 0.2, Inches(0.11), Inches(0.11),
             fill=SILVER, shape=MSO_SHAPE.OVAL)
        tri = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,
                                     icx + ic * 0.30, icy + ic * 0.44,
                                     Inches(0.30), Inches(0.18))
        tri.fill.solid(); tri.fill.fore_color.rgb = SILVER
        tri.line.fill.background(); tri.shadow.inherit = False
        ly = icy + ic + Inches(0.10)
    else:
        ly = y + h / 2 - Inches(0.22)
    textbox(slide, x, ly, w, Inches(0.4),
            [{"text": label, "size": 10.5, "color": SLATE, "bold": True,
              "spacing": 2.0, "align": PP_ALIGN.CENTER}], align=PP_ALIGN.CENTER)
    if note:
        textbox(slide, x, ly + Inches(0.30), w, Inches(0.3),
                [{"text": note, "size": 8, "color": PLATINUM, "spacing": 0.8,
                  "align": PP_ALIGN.CENTER}], align=PP_ALIGN.CENTER)


def icon_chip(slide, x, y, size=Inches(0.6)):
    rect(slide, x, y, size, size, fill=None, line=CHRYSLER, line_w=Pt(1.25),
         shape=MSO_SHAPE.OVAL)
    inner = Inches(0.18)
    rect(slide, x + (size - inner) / 2, y + (size - inner) / 2, inner, inner,
         fill=CHRYSLER, shape=MSO_SHAPE.OVAL)


def wing_mark(slide, x, y, w=Inches(2.0)):
    hairline(slide, x, y, w, color=SILVER, weight=Pt(1.5))
    d = Inches(0.14)
    rect(slide, x + w / 2 - d / 2, y - d / 2, d, d, fill=CHRYSLER,
         shape=MSO_SHAPE.DIAMOND)


def callout(slide, x, y, w, h, head, sub, num=None):
    rect(slide, x, y, w, h, fill=CLOUD, line=PLATINUM, line_w=Pt(0.5))
    rect(slide, x, y, Inches(0.06), h, fill=CHRYSLER)
    lines = []
    if num:
        lines.append({"text": num, "size": 11, "color": SILVER, "bold": True,
                      "spacing": 1.5, "space_after": 2})
    lines += [{"text": head, "size": 13, "color": INK, "bold": True,
               "space_after": 3},
              {"text": sub, "size": 10.5, "color": SLATE, "line_spacing": 1.2}]
    textbox(slide, x + Inches(0.3), y + Inches(0.18), w - Inches(0.5),
            h - Inches(0.32), lines)


def col_x(i, n, x0=Inches(0.75), total=Inches(11.83), gap=Inches(0.4)):
    w = Emu(int((total - gap * (n - 1)) / n))
    return Emu(int(x0 + (w + gap) * i)), w


# ----------------------------------------------------------------------------
# Opening / keynote templates
# ----------------------------------------------------------------------------
def s_title():
    s = add_slide(); set_bg(s)
    image_frame(s, Inches(0), Inches(0), SLIDE_W, SLIDE_H,
                label="FULL-WIDTH HERO — PORTFOLIO KEY VISUAL", fill=CLOUD,
                note="Replace · all six concepts together")
    rect(s, Inches(0), Inches(4.55), Inches(9.0), Inches(2.95), fill=WHITE)
    wing_mark(s, Inches(0.85), Inches(5.0), Inches(2.0))
    kicker(s, Inches(0.85), Inches(5.2), "Chrysler · CES")
    textbox(s, Inches(0.82), Inches(5.55), Inches(8.0), Inches(0.9),
            [{"text": "Chrysler at CES", "size": 44, "color": INK, "bold": True,
              "name": FONT_DISPLAY}])
    textbox(s, Inches(0.85), Inches(6.5), Inches(8.0), Inches(0.4),
            [{"text": "Six concepts. One portfolio. One story.", "size": 16,
              "color": GRAPHITE}])
    textbox(s, Inches(0.85), Inches(6.95), Inches(11.5), Inches(0.4),
            [{"text": "Chrysler 300 · Pacifica Pinnacle · Pacifica Grizzly Peak "
                      "· C2U · C2X · Airflow", "size": 10.5, "color": SLATE,
              "spacing": 0.5}])
    notes(s, "Open the keynote: every concept on the CES floor is one portfolio "
             "telling one story — the rebirth of Chrysler through technology, "
             "simplicity, and attainability. Set an aspirational, confident tone.")
    footer(s, page())
    return s


def s_portfolio():
    s = add_slide(); set_bg(s)
    kicker(s, Inches(0.75), Inches(0.55), "The Portfolio")
    textbox(s, Inches(0.72), Inches(0.9), Inches(11.5), Inches(0.7),
            [{"text": "One Portfolio at CES", "size": 26, "color": INK,
              "bold": True, "name": FONT_DISPLAY}])
    textbox(s, Inches(0.75), Inches(1.6), Inches(11.5), Inches(0.4),
            [{"text": "Six concepts, designed to complement one another and add "
                      "up to a single Chrysler story.", "size": 13,
              "color": GRAPHITE}])
    names = ["Chrysler 300", "Pacifica Pinnacle", "Pacifica Grizzly Peak",
             "C2U", "C2X", "Airflow"]
    for i, name in enumerate(names):
        r, c = divmod(i, 3)
        x, w = col_x(c, 3)
        y = Inches(2.25) + Inches(2.35) * r
        image_frame(s, x, y, w, Inches(1.95), label=f"CONCEPT {i+1}",
                    note="Replace", icon=(r == 0))
        textbox(s, x, Emu(int(y + Inches(2.0))), w, Inches(0.3),
                [{"text": name, "size": 12, "color": INK, "bold": True,
                  "align": PP_ALIGN.CENTER}], align=PP_ALIGN.CENTER)
    notes(s, "Reveal the full portfolio at once. Emphasize breadth — flagship, "
             "family, accessible, electric — all unmistakably Chrysler. Each is "
             "expanded later in the deck.")
    footer(s, page())
    return s


def s_statement(kick, title, statement, supports=None):
    """Keynote statement slide: bold idea + optional three supports + hero."""
    s = add_slide(); set_bg(s)
    kicker(s, Inches(0.75), Inches(0.95), kick)
    textbox(s, Inches(0.72), Inches(1.35), Inches(5.9), Inches(1.4),
            [{"text": title, "size": 30, "color": INK, "bold": True,
              "name": FONT_DISPLAY, "line_spacing": 1.0}])
    hairline(s, Inches(0.75), Inches(2.85), Inches(5.6))
    textbox(s, Inches(0.75), Inches(3.05), Inches(5.7), Inches(1.6),
            [{"text": statement, "size": 16, "color": GRAPHITE,
              "line_spacing": 1.3}])
    if supports:
        by = Inches(4.55)
        for i, (h, sub) in enumerate(supports):
            y = by + Inches(0.82) * i
            icon_chip(s, Inches(0.78), y, Inches(0.42))
            textbox(s, Inches(1.4), y - Inches(0.04), Inches(5.0), Inches(0.8),
                    [{"text": h, "size": 13, "color": INK, "bold": True,
                      "space_after": 1},
                     {"text": sub, "size": 10.5, "color": SLATE}])
    image_frame(s, Inches(7.1), Inches(0.95), Inches(5.48), Inches(5.5),
                label="SUPPORTING IMAGE")
    notes(s, f"{title}: deliver one clear idea. Keep the statement short and "
             "let the image carry the feeling. This frames the concepts that "
             "follow.")
    footer(s, page())
    return s


def s_beat(idx, title, statement):
    """One keynote story beat — large, cinematic, one idea."""
    s = add_slide(); set_bg(s)
    image_frame(s, Inches(6.4), Inches(0), Inches(6.93), SLIDE_H,
                label="BEAT IMAGE", fill=CLOUD, note="Replace · story visual")
    rect(s, Inches(0), Inches(0), Inches(6.6), SLIDE_H, fill=WHITE)
    kicker(s, Inches(0.75), Inches(2.3), f"Keynote · Story {idx:02d}")
    textbox(s, Inches(0.72), Inches(2.75), Inches(5.6), Inches(2.0),
            [{"text": title, "size": 34, "color": INK, "bold": True,
              "name": FONT_DISPLAY, "line_spacing": 1.0}])
    hairline(s, Inches(0.75), Inches(4.7), Inches(3.5))
    textbox(s, Inches(0.75), Inches(4.9), Inches(5.4), Inches(1.4),
            [{"text": statement, "size": 15, "color": GRAPHITE,
              "line_spacing": 1.3}])
    notes(s, f"Keynote beat {idx}: '{title}'. One cinematic idea, one line, one "
             "image. Pause here — this is a story moment, not a data slide.")
    footer(s, page())
    return s


def s_roadmap():
    s = add_slide(); set_bg(s)
    kicker(s, Inches(0.75), Inches(0.55), "Story & Roadmap")
    textbox(s, Inches(0.72), Inches(0.9), Inches(11.5), Inches(0.7),
            [{"text": "Concept Story Meets Technology Roadmap", "size": 24,
              "color": INK, "bold": True, "name": FONT_DISPLAY}])
    textbox(s, Inches(0.75), Inches(1.6), Inches(11.5), Inches(0.4),
            [{"text": "How the exterior and interior story lines up with the "
                      "technology roadmap.", "size": 13, "color": GRAPHITE}])
    # exterior / interior story rails
    image_frame(s, Inches(0.75), Inches(2.2), Inches(5.83), Inches(1.7),
                label="EXTERIOR STORY", note="Replace", icon=False)
    image_frame(s, Inches(0.75), Inches(4.05), Inches(5.83), Inches(1.7),
                label="INTERIOR STORY", note="Replace", icon=False)
    # roadmap timeline (right)
    rx, rw = Inches(7.0), Inches(5.58)
    hairline(s, rx, Inches(4.0), rw, color=SILVER, weight=Pt(1.5))
    phases = ["Today", "Near-Term", "Mid-Term", "Vision"]
    for i, ph in enumerate(phases):
        x = Emu(int(rx + (rw / (len(phases) - 1)) * i))
        rect(s, Emu(int(x - Inches(0.08))), Inches(3.92), Inches(0.16),
             Inches(0.16), fill=CHRYSLER, shape=MSO_SHAPE.OVAL)
        textbox(s, Emu(int(x - Inches(0.7))), Inches(3.35), Inches(1.4),
                Inches(0.3),
                [{"text": ph, "size": 11, "color": INK, "bold": True,
                  "align": PP_ALIGN.CENTER}], align=PP_ALIGN.CENTER)
        textbox(s, Emu(int(x - Inches(0.7))), Inches(4.25), Inches(1.4),
                Inches(1.0),
                [{"text": "[ milestone ]", "size": 9, "color": SLATE,
                  "align": PP_ALIGN.CENTER, "line_spacing": 1.1}],
                align=PP_ALIGN.CENTER)
    textbox(s, Inches(7.0), Inches(2.35), Inches(5.4), Inches(0.8),
            [{"text": "Each concept plots against the roadmap — showing how "
                      "design intent and technology maturity advance together.",
              "size": 12, "color": GRAPHITE, "line_spacing": 1.25}])
    notes(s, "Tie the design narrative to the technology roadmap: the exterior "
             "and interior stories are not styling exercises — they visualize "
             "where the technology is heading and when. Plot each concept along "
             "Today → Vision.")
    footer(s, page())
    return s


# ----------------------------------------------------------------------------
# Per-concept templates
# ----------------------------------------------------------------------------
def s_concept_intro(idx, name, role, beat):
    s = add_slide(); set_bg(s)
    image_frame(s, Inches(5.0), Inches(0), Inches(8.333), SLIDE_H,
                label="CONCEPT HERO", fill=CLOUD, note="Replace · signature shot")
    rect(s, Inches(0), Inches(0), Inches(5.0), SLIDE_H, fill=WHITE)
    wing_mark(s, Inches(0.75), Inches(1.6), Inches(1.8))
    kicker(s, Inches(0.75), Inches(1.8), f"Concept {idx:02d}")
    textbox(s, Inches(0.72), Inches(2.25), Inches(4.1), Inches(1.6),
            [{"text": name, "size": 34, "color": INK, "bold": True,
              "name": FONT_DISPLAY, "line_spacing": 1.0}])
    textbox(s, Inches(0.75), Inches(3.9), Inches(3.95), Inches(1.4),
            [{"text": role, "size": 14, "color": GRAPHITE, "line_spacing": 1.3}])
    hairline(s, Inches(0.75), Inches(5.7), Inches(3.5))
    textbox(s, Inches(0.75), Inches(5.85), Inches(4.0), Inches(0.7),
            [{"text": "STORY", "size": 9, "color": SLATE, "bold": True,
              "spacing": 2.0, "space_after": 2},
             {"text": beat, "size": 12, "color": CHRYSLER, "bold": True}])
    notes(s, f"{name}: open the concept. State its role in the portfolio and the "
             f"keynote beat it carries ({beat}). One confident line — let the "
             "hero image do the work.")
    footer(s, page())
    return s


def s_design(kick, title, caption):
    """Full-width exterior or interior image + one caption."""
    s = add_slide(); set_bg(s)
    kicker(s, Inches(0.75), Inches(0.6), kick)
    textbox(s, Inches(0.72), Inches(0.95), Inches(11.0), Inches(0.7),
            [{"text": title, "size": 24, "color": INK, "bold": True,
              "name": FONT_DISPLAY}])
    image_frame(s, Inches(0.75), Inches(1.8), Inches(11.83), Inches(4.65),
                label=f"LARGE {title.split('—')[-1].strip().upper()} IMAGE")
    textbox(s, Inches(0.75), Inches(6.55), Inches(11.0), Inches(0.4),
            [{"text": caption, "size": 12, "color": SLATE, "spacing": 0.5}])
    notes(s, f"{title}: one immersive image, one short caption. Speak to the "
             "design intent and the feeling, not the spec sheet.")
    footer(s, page())
    return s


def s_innovations(kick, title, hint, items=None):
    """Large image + three innovation callouts (populated or placeholder)."""
    s = add_slide(); set_bg(s)
    kicker(s, Inches(0.75), Inches(0.6), kick)
    textbox(s, Inches(0.72), Inches(0.95), Inches(11.0), Inches(0.7),
            [{"text": title, "size": 24, "color": INK, "bold": True,
              "name": FONT_DISPLAY}])
    image_frame(s, Inches(0.75), Inches(1.85), Inches(7.0), Inches(4.6),
                label="INNOVATION IMAGE", note="Replace · call out details")
    for i in range(3):
        y = Emu(int(Inches(1.85) + Inches(1.6) * i))
        if items:
            head, sub = items[i]
        else:
            head, sub = "[ Innovation headline ]", hint
        callout(s, Inches(8.0), y, Inches(4.58), Inches(1.42),
                head, sub, num=f"0{i+1}")
    notes(s, f"{title}: highlight three innovations. For each, name it plainly "
             "and say how it simplifies the vehicle, removes complexity, or makes "
             "the experience more attainable — the through-line of the keynote.")
    footer(s, page())
    return s


def s_acm(kick, name):
    """Advanced Color & Materials (ACM) — swatch rail + material mood."""
    s = add_slide(); set_bg(s)
    kicker(s, Inches(0.75), Inches(0.55), kick)
    textbox(s, Inches(0.72), Inches(0.9), Inches(7.5), Inches(0.7),
            [{"text": "Advanced Color & Materials", "size": 24, "color": INK,
              "bold": True, "name": FONT_DISPLAY}])
    textbox(s, Inches(0.75), Inches(1.65), Inches(6.0), Inches(1.1),
            [{"text": f"The color, material, and finish story for {name} — warm, "
                      "responsible, and quietly premium.", "size": 14,
              "color": GRAPHITE, "line_spacing": 1.3}])
    image_frame(s, Inches(7.1), Inches(0.95), Inches(5.48), Inches(5.5),
                label="MATERIAL MOOD", note="Replace · CMF / trim mood")
    swatches = [("Warm Grey", RGBColor(0xB9, 0xB4, 0xAB)),
                ("Recycled Slate", CHRYSLER),
                ("Natural Linen", RGBColor(0xE7, 0xE2, 0xD8)),
                ("Graphite", GRAPHITE),
                ("Pearl", RGBColor(0xF2, 0xF3, 0xF4))]
    sy, sw, sh, sgap = Inches(3.15), Inches(1.15), Inches(1.15), Inches(0.18)
    for i, (nm, col) in enumerate(swatches):
        x = Emu(int(Inches(0.75) + (sw + sgap) * i))
        rect(s, x, sy, sw, sh, fill=col, line=PLATINUM, line_w=Pt(0.5))
        textbox(s, x, Emu(int(sy + sh + Inches(0.1))), sw, Inches(0.4),
                [{"text": nm, "size": 8.5, "color": SLATE, "spacing": 0.5}])
    textbox(s, Inches(0.75), Inches(5.4), Inches(6.0), Inches(0.5),
            [{"text": "Swatches are placeholders — replace with the approved "
                      "ACM palette and trim samples.", "size": 9.5,
              "color": PLATINUM}])
    notes(s, f"{name} ACM: tell the material story — recycled and traceable "
             "content, warm tactility, and finishes that read premium through "
             "restraint. Keep it about feeling and responsibility, not specs.")
    footer(s, page())
    return s


def concept_block(idx, name, role, beat, ext_cap, int_cap,
                  ext_innov=None, int_innov=None):
    s_concept_intro(idx, name, role, beat)
    s_design(f"{name} · Exterior", f"{name} — Exterior", ext_cap)
    s_innovations(f"{name} · Exterior", f"{name} — Exterior Innovations",
                  "[ how it simplifies or removes complexity ]", ext_innov)
    s_design(f"{name} · Interior", f"{name} — Interior", int_cap)
    s_innovations(f"{name} · Interior", f"{name} — Interior Innovations",
                  "[ how it eases everyday use ]", int_innov)
    s_acm(f"{name} · ACM", name)


# ----------------------------------------------------------------------------
# Brand foundation (carried over from the Chrysler Brand Design Vision deck)
# ----------------------------------------------------------------------------
def s_why():
    s = add_slide(); set_bg(s)
    kicker(s, Inches(0.75), Inches(0.95), "The Vision")
    textbox(s, Inches(0.72), Inches(1.35), Inches(5.9), Inches(1.0),
            [{"text": "Why Chrysler?", "size": 32, "color": INK, "bold": True,
              "name": FONT_DISPLAY}])
    hairline(s, Inches(0.75), Inches(2.5), Inches(5.6))
    textbox(s, Inches(0.75), Inches(2.75), Inches(5.7), Inches(2.4),
            [{"text": "Chrysler removes friction from everyday life through "
                      "thoughtful, people-first design.", "size": 22,
              "color": GRAPHITE, "line_spacing": 1.25}])
    textbox(s, Inches(0.75), Inches(5.4), Inches(5.7), Inches(1.0),
            [{"text": "Every concept at CES is evidence of this one idea.",
              "size": 13, "color": CHRYSLER, "italic": True}])
    image_frame(s, Inches(7.1), Inches(0.95), Inches(5.48), Inches(5.5),
                label="BRAND VISION IMAGE")
    notes(s, "Ground the CES story in the brand vision: the single philosophy "
             "every concept serves. This is the same idea that opens the Brand "
             "Design Vision deck — keep the two consistent.")
    footer(s, page())
    return s


def s_pillars():
    s = add_slide(); set_bg(s)
    kicker(s, Inches(0.75), Inches(0.55), "Design Principles")
    textbox(s, Inches(0.72), Inches(0.9), Inches(11.5), Inches(0.7),
            [{"text": "The Four Pillars", "size": 26, "color": INK,
              "bold": True, "name": FONT_DISPLAY}])
    textbox(s, Inches(0.75), Inches(1.6), Inches(11.5), Inches(0.4),
            [{"text": "The principles behind every Chrysler design decision.",
              "size": 13, "color": GRAPHITE}])
    pillars = [
        ("People First", "Every decision begins with the people who use it."),
        ("Everyday Ingenuity", "Thoughtful solutions that make daily life easier."),
        ("Human-Centered Intelligence", "Technology that quietly supports."),
        ("Modern American Design", "Confident, optimistic, unmistakably Chrysler."),
    ]
    for i, (h, sub) in enumerate(pillars):
        x, w = col_x(i, 4)
        icon_chip(s, x, Inches(2.6), Inches(0.5))
        hairline(s, x, Inches(3.4), Emu(int(w)))
        textbox(s, x, Inches(3.55), w, Inches(2.6),
                [{"text": f"0{i+1}", "size": 12, "color": SILVER, "bold": True,
                  "spacing": 2.0, "space_after": 6},
                 {"text": h, "size": 15, "color": INK, "bold": True,
                  "space_after": 5},
                 {"text": sub, "size": 11, "color": SLATE, "line_spacing": 1.2}])
    notes(s, "The four pillars carried over from the brand vision. At CES they "
             "become the lens for reading every concept and every innovation.")
    footer(s, page())
    return s


# ----------------------------------------------------------------------------
# Build
# ----------------------------------------------------------------------------
# Opening & keynote narrative
s_title()
s_portfolio()
s_why()        # brand foundation carried over from the Vision deck
s_pillars()    # brand foundation carried over from the Vision deck
s_statement(
    "The Through-Line", "One Portfolio, One Story",
    "Six concepts that complement one another — together they tell the rebirth "
    "of Chrysler through technology.",
    [("They share a philosophy", "One design language across every concept."),
     ("They share a purpose", "Technology in service of simplicity."),
     ("They share a story", "Each advances the same keynote arc.")])
s_beat(1, "The Rebirth of Chrysler",
       "A confident return — Chrysler reintroduced through design and intent.")
s_beat(2, "Chrysler Is About Technology",
       "Technology is the core of the brand — the engine of everything we show.")
s_beat(3, "Technology Through Simplicity & Affordability",
       "Advanced technology, made simple and attainable — not complex or "
       "exclusive.")
s_beat(4, "Technology That Removes Complexity",
       "Technology lets us remove things from the vehicle — making it simpler, "
       "cleaner, and more attainable.")
s_statement(
    "The Customer", "Why Our Customer Cares",
    "Features only matter when they make life easier — we lead with the benefit, "
    "not the technology.",
    [("Less to manage", "Complexity removed, not added."),
     ("More attainable", "Premium experience, accessible price."),
     ("Designed for real life", "Technology that quietly serves people.")])
s_roadmap()

# Concept expansions (intro / ext / ext innovations / int / int innovations / ACM)
concept_block(
    1, "Chrysler 300",
    "The flagship expression of Chrysler design — composed, commanding, never "
    "excessive.",
    "Rebirth of Chrysler",
    "Commanding, composed proportion — presence without excess.",
    "An interior sanctuary — crafted, composed, human in scale.",
    ext_innov=[("Commanding Presence", "Architectural proportion, balanced and tailored."),
               ("Quiet Power", "Confidence expressed through restraint."),
               ("Timeless Craftsmanship", "Authentic American design, refined.")],
    int_innov=[("Crafted Sanctuary", "Calm materials, composed and human in scale."),
               ("Quiet Technology", "Intelligence that recedes into the cabin."),
               ("Tailored Detail", "Restraint expressed through craft.")])
concept_block(
    2, "Pacifica Pinnacle",
    "The evolution of the modern family vehicle — flexible, sophisticated, and "
    "designed around people.",
    "Technology through simplicity",
    "Refined, confident family form — familiar, elevated.",
    "A family sanctuary — flexible, comfortable, designed around people.",
    ext_innov=[("Flexible by Design", "Space that adapts to every journey."),
               ("Sophisticated Comfort", "Refined, calm, effortlessly livable."),
               ("Technology Around People", "Intelligence that serves the family.")],
    int_innov=[("Serene Lounge", "Comfort and control, made simple."),
               ("Adaptive Space", "Flexible seating for real family life."),
               ("Effortless Control", "Technology that quietly serves.")])
concept_block(
    3, "Pacifica Grizzly Peak",
    "Capable family adventure — attainable capability, thoughtfully packaged.",
    "Simpler and more attainable",
    "Rugged, ready stance — capability with restraint.",
    "A durable, flexible cabin — built for real life.")
concept_block(
    4, "C2U",
    "[ Concept role placeholder ] — accessible, urban, and design-led.",
    "Simplicity & affordability",
    "Compact, confident urban form — approachable design.",
    "A smart, uncluttered cabin — attainable by design.")
concept_block(
    5, "C2X",
    "[ Concept role placeholder ] — versatile, attainable, and modern.",
    "Simplicity & affordability",
    "Versatile crossover form — capable and clean.",
    "A flexible, thoughtful interior — space made simple.")
concept_block(
    6, "Airflow",
    "Airflow introduces Chrysler's philosophy: innovative practicality, "
    "designed around people.",
    "Chrysler is about technology",
    "Aerodynamic, future-forward silhouette — technology as form.",
    "Designed around people — flexible, calm, effortlessly useful.",
    ext_innov=[("Spacious & Adaptive", "Versatile space for everyday life."),
               ("Technology With Purpose", "Intelligence that serves, then disappears."),
               ("Confident Simplicity", "Clean, logical, human-centered form.")],
    int_innov=[("Calm & Connected", "An intelligent cabin that stays quiet."),
               ("Effortless Interface", "Technology that recedes into use."),
               ("Human-Centered Space", "Designed around people, not features.")])

OUT = "/home/user/CLAUDE/Chrysler_CES_Concept_Portfolio.pptx"
prs.save(OUT)
print(f"Saved {OUT} with {len(prs.slides._sldIdLst)} slides.")
