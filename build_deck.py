#!/usr/bin/env python3
"""
Stellantis Executive Portfolio Deck — scaffold generator.

Produces a premium, minimal, automotive executive deck (16:9) with placeholder
slides only: titles, layouts, image placeholders, icon placeholders, and text
boxes for the design team to populate.

Design language: minimal / premium / modern / automotive. White backgrounds,
large imagery, very little text, restrained Chrysler-inspired typography.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ----------------------------------------------------------------------------
# Design system
# ----------------------------------------------------------------------------
EMU_PER_INCH = 914400
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

# Palette — restrained, premium automotive
INK        = RGBColor(0x1A, 0x1A, 0x1C)   # near-black, primary text
GRAPHITE   = RGBColor(0x4A, 0x4E, 0x57)   # secondary text
SLATE      = RGBColor(0x86, 0x8B, 0x94)   # tertiary / captions
PLATINUM   = RGBColor(0xC8, 0xCC, 0xD2)   # hairlines / accents
MIST       = RGBColor(0xEF, 0xF1, 0xF3)   # placeholder fill
CLOUD      = RGBColor(0xF7, 0xF8, 0xF9)   # subtle panel
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
CHRYSLER   = RGBColor(0x1B, 0x2A, 0x4A)   # deep Chrysler-inspired navy accent
SILVER     = RGBColor(0xA9, 0xAE, 0xB5)   # wing/badge silver

# Typography — system-safe premium pairing
FONT_DISPLAY = "Arial"          # substitute for a licensed Chrysler display face
FONT_BODY    = "Arial"

# ----------------------------------------------------------------------------
# Helpers
# ----------------------------------------------------------------------------
prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
BLANK = prs.slide_layouts[6]


def add_slide():
    return prs.slides.add_slide(BLANK)


def set_bg(slide, color=WHITE):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _set_font(run, size, color, bold=False, name=FONT_BODY, spacing=None):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.name = name
    f.color.rgb = color
    if spacing is not None:
        # letter spacing in 1/100 pt via XML
        rPr = run._r.get_or_add_rPr()
        rPr.set("spc", str(int(spacing * 100)))


def textbox(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """lines: list of dicts {text,size,color,bold,name,spacing,space_after}"""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ln.get("align", align)
        if ln.get("space_after") is not None:
            p.space_after = Pt(ln["space_after"])
        if ln.get("space_before") is not None:
            p.space_before = Pt(ln["space_before"])
        if ln.get("line_spacing") is not None:
            p.line_spacing = ln["line_spacing"]
        r = p.add_run()
        r.text = ln["text"]
        _set_font(
            r,
            ln.get("size", 18),
            ln.get("color", INK),
            ln.get("bold", False),
            ln.get("name", FONT_BODY),
            ln.get("spacing"),
        )
    return tb


def rect(slide, x, y, w, h, fill=None, line=None, line_w=None, shape=MSO_SHAPE.RECTANGLE):
    sp = slide.shapes.add_shape(shape, x, y, w, h)
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = line_w or Pt(0.75)
    return sp


def hairline(slide, x, y, w, color=PLATINUM, weight=Pt(1.0)):
    ln = slide.shapes.add_connector(2, x, y, x + w, y)
    ln.line.color.rgb = color
    ln.line.width = weight
    return ln


def kicker(slide, x, y, text, color=CHRYSLER):
    """Small uppercase eyebrow label with letter spacing."""
    return textbox(
        slide, x, y, Inches(6), Inches(0.3),
        [{"text": text.upper(), "size": 11, "color": color, "bold": True,
          "name": FONT_DISPLAY, "spacing": 3.0}],
    )


def page_meta(slide, number, section="STELLANTIS  ·  DESIGN PORTFOLIO"):
    """Footer: confidential label + page number."""
    textbox(slide, Inches(0.7), Inches(7.02), Inches(8), Inches(0.3),
            [{"text": section, "size": 8, "color": SLATE, "spacing": 1.5}])
    textbox(slide, Inches(11.6), Inches(7.02), Inches(1.0), Inches(0.3),
            [{"text": f"{number:02d}", "size": 8, "color": SLATE, "spacing": 1.5}],
            align=PP_ALIGN.RIGHT)
    textbox(slide, Inches(11.6), Inches(0.32), Inches(1.0), Inches(0.3),
            [{"text": "CONFIDENTIAL", "size": 7.5, "color": PLATINUM, "spacing": 2.0}],
            align=PP_ALIGN.RIGHT)


def image_placeholder(slide, x, y, w, h, label="HERO IMAGE", fill=MIST,
                      note="Replace · 16:9 · 300dpi"):
    """A premium image placeholder: soft panel, centered crosshair icon + label."""
    panel = rect(slide, x, y, w, h, fill=fill)
    # corner ticks for a 'crop frame' premium feel
    tick = Inches(0.28)
    tw = Pt(1.25)
    corners = [
        (x, y, "tl"), (x + w, y, "tr"),
        (x, y + h, "bl"), (x + w, y + h, "br"),
    ]
    for cx, cy, pos in corners:
        dx = tick if "l" in pos else -tick
        dy = tick if "t" in pos else -tick
        h1 = slide.shapes.add_connector(2, cx, cy, cx + dx, cy)
        h1.line.color.rgb = SILVER
        h1.line.width = tw
        v1 = slide.shapes.add_connector(2, cx, cy, cx, cy + dy)
        v1.line.color.rgb = SILVER
        v1.line.width = tw
    # center icon — simple framed "image" glyph
    ic = Inches(0.62)
    icx = x + (w - ic) / 2
    icy = y + (h - ic) / 2 - Inches(0.22)
    glyph = rect(slide, icx, icy, ic, ic, fill=None, line=SILVER, line_w=Pt(1.5),
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    # mountain + sun motif inside glyph
    sun = rect(slide, icx + ic * 0.22, icy + ic * 0.2, Inches(0.12), Inches(0.12),
               fill=SILVER, shape=MSO_SHAPE.OVAL)
    tri = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,
                                 icx + ic * 0.30, icy + ic * 0.42,
                                 Inches(0.34), Inches(0.22))
    tri.fill.solid(); tri.fill.fore_color.rgb = SILVER; tri.line.fill.background()
    tri.shadow.inherit = False
    # label under icon
    textbox(slide, x, icy + ic + Inches(0.12), w, Inches(0.4),
            [{"text": label, "size": 11, "color": SLATE, "bold": True,
              "spacing": 2.5, "align": PP_ALIGN.CENTER}], align=PP_ALIGN.CENTER)
    textbox(slide, x, icy + ic + Inches(0.42), w, Inches(0.3),
            [{"text": note, "size": 8.5, "color": PLATINUM,
              "spacing": 1.0, "align": PP_ALIGN.CENTER}], align=PP_ALIGN.CENTER)
    return panel


def icon_chip(slide, x, y, size=Inches(0.5)):
    """Small circular icon placeholder for bullets / features."""
    c = rect(slide, x, y, size, size, fill=None, line=CHRYSLER, line_w=Pt(1.25),
             shape=MSO_SHAPE.OVAL)
    inner = Inches(0.16)
    rect(slide, x + (size - inner) / 2, y + (size - inner) / 2, inner, inner,
         fill=CHRYSLER, shape=MSO_SHAPE.OVAL)
    return c


def wing_mark(slide, x, y, w=Inches(2.4)):
    """Abstract Chrysler 'wing' inspired hairline mark (placeholder for badge)."""
    cy = y
    hairline(slide, x, cy, w, color=SILVER, weight=Pt(1.5))
    # small diamond at center
    d = Inches(0.14)
    rect(slide, x + w / 2 - d / 2, cy - d / 2, d, d, fill=CHRYSLER,
         shape=MSO_SHAPE.DIAMOND)


# ----------------------------------------------------------------------------
# Slide templates
# ----------------------------------------------------------------------------
def slide_title(number, kick, title, subtitle, footnote):
    s = add_slide(); set_bg(s)
    # full-bleed hero band on the right two-thirds
    image_placeholder(s, Inches(5.0), Inches(0), Inches(8.333), SLIDE_H,
                      label="FULL-BLEED HERO", fill=CLOUD,
                      note="Replace · cinematic vehicle key visual")
    # left editorial column on white
    rect(s, Inches(0), Inches(0), Inches(5.0), SLIDE_H, fill=WHITE)
    wing_mark(s, Inches(0.75), Inches(1.5), Inches(2.0))
    kicker(s, Inches(0.75), Inches(1.75), kick)
    textbox(s, Inches(0.72), Inches(2.2), Inches(4.0), Inches(2.6),
            [{"text": title, "size": 40, "color": INK, "bold": True,
              "name": FONT_DISPLAY, "line_spacing": 1.0, "space_after": 10}])
    textbox(s, Inches(0.75), Inches(4.5), Inches(3.9), Inches(1.2),
            [{"text": subtitle, "size": 14, "color": GRAPHITE,
              "line_spacing": 1.25}])
    hairline(s, Inches(0.75), Inches(6.4), Inches(3.5))
    textbox(s, Inches(0.75), Inches(6.55), Inches(4.0), Inches(0.6),
            [{"text": footnote, "size": 9.5, "color": SLATE, "spacing": 1.0}])
    return s


def slide_strategy(number, kick, title, summary, bullets):
    """Strategy layout: summary sentence + 3 bullets (left), hero image (right)."""
    s = add_slide(); set_bg(s)
    # right hero
    image_placeholder(s, Inches(7.1), Inches(0.9), Inches(5.55), Inches(5.5),
                      label="HERO IMAGE")
    # left content column
    kicker(s, Inches(0.75), Inches(0.95), kick)
    textbox(s, Inches(0.72), Inches(1.35), Inches(5.9), Inches(1.1),
            [{"text": title, "size": 30, "color": INK, "bold": True,
              "name": FONT_DISPLAY, "line_spacing": 1.0}])
    hairline(s, Inches(0.75), Inches(2.35), Inches(5.6))
    # executive summary sentence
    textbox(s, Inches(0.75), Inches(2.6), Inches(5.7), Inches(1.2),
            [{"text": summary, "size": 16, "color": GRAPHITE,
              "line_spacing": 1.3}])
    # three bullets with icon chips
    by = Inches(4.15)
    step = Inches(0.92)
    for i, (head, sub) in enumerate(bullets):
        y = by + step * i
        icon_chip(s, Inches(0.78), y + Inches(0.02))
        textbox(s, Inches(1.5), y - Inches(0.05), Inches(5.0), Inches(0.9),
                [{"text": head, "size": 14, "color": INK, "bold": True,
                  "space_after": 2},
                 {"text": sub, "size": 11.5, "color": SLATE,
                  "line_spacing": 1.15}])
    page_meta(s, number)
    return s


def slide_design(number, kick, title, caption, layout="single"):
    """Design layout: imagery-forward with one short caption."""
    s = add_slide(); set_bg(s)
    kicker(s, Inches(0.75), Inches(0.55), kick)
    textbox(s, Inches(0.72), Inches(0.9), Inches(9.0), Inches(0.7),
            [{"text": title, "size": 24, "color": INK, "bold": True,
              "name": FONT_DISPLAY}])
    if layout == "single":
        image_placeholder(s, Inches(0.75), Inches(1.75), Inches(11.83), Inches(4.7),
                          label="FULL-WIDTH DESIGN IMAGE",
                          note="Replace · hero exterior / interior render")
    elif layout == "duo":
        image_placeholder(s, Inches(0.75), Inches(1.75), Inches(5.83), Inches(4.7),
                          label="PRIMARY VIEW")
        image_placeholder(s, Inches(6.75), Inches(1.75), Inches(5.83), Inches(4.7),
                          label="SECONDARY VIEW")
    elif layout == "triptych":
        image_placeholder(s, Inches(0.75), Inches(1.75), Inches(7.4), Inches(4.7),
                          label="HERO DETAIL")
        image_placeholder(s, Inches(8.35), Inches(1.75), Inches(4.23), Inches(2.27),
                          label="DETAIL A", note="Replace")
        image_placeholder(s, Inches(8.35), Inches(4.18), Inches(4.23), Inches(2.27),
                          label="DETAIL B", note="Replace")
    # single short caption
    textbox(s, Inches(0.75), Inches(6.55), Inches(11.0), Inches(0.4),
            [{"text": caption, "size": 12, "color": SLATE, "spacing": 0.5}])
    page_meta(s, number)
    return s


def slide_portfolio(number, kick, title, summary, tiles):
    """Portfolio grid: summary + row of vehicle tiles with captions."""
    s = add_slide(); set_bg(s)
    kicker(s, Inches(0.75), Inches(0.55), kick)
    textbox(s, Inches(0.72), Inches(0.9), Inches(11.0), Inches(0.7),
            [{"text": title, "size": 24, "color": INK, "bold": True,
              "name": FONT_DISPLAY}])
    textbox(s, Inches(0.75), Inches(1.6), Inches(11.5), Inches(0.6),
            [{"text": summary, "size": 14, "color": GRAPHITE, "line_spacing": 1.25}])
    n = len(tiles)
    gap = Inches(0.4)
    total = Inches(11.83)
    tw = Emu(int((total - gap * (n - 1)) / n))
    x0 = Inches(0.75)
    for i, (name, cap) in enumerate(tiles):
        x = Emu(int(x0 + (tw + gap) * i))
        image_placeholder(s, x, Inches(2.5), tw, Inches(3.2),
                          label=f"CONCEPT {i+1}", note="Replace")
        textbox(s, x, Inches(5.85), tw, Inches(0.4),
                [{"text": name, "size": 13, "color": INK, "bold": True}])
        textbox(s, x, Inches(6.2), tw, Inches(0.5),
                [{"text": cap, "size": 10, "color": SLATE, "line_spacing": 1.1}])
    page_meta(s, number)
    return s


def slide_cmf(number, kick, title, summary, swatches):
    """Color / Material / Finish strategy — swatch rail + material image."""
    s = add_slide(); set_bg(s)
    kicker(s, Inches(0.75), Inches(0.55), kick)
    textbox(s, Inches(0.72), Inches(0.9), Inches(7.0), Inches(0.7),
            [{"text": title, "size": 24, "color": INK, "bold": True,
              "name": FONT_DISPLAY}])
    textbox(s, Inches(0.75), Inches(1.65), Inches(6.0), Inches(1.3),
            [{"text": summary, "size": 15, "color": GRAPHITE, "line_spacing": 1.3}])
    # material mood image right
    image_placeholder(s, Inches(7.1), Inches(0.95), Inches(5.55), Inches(5.5),
                      label="MATERIAL MOOD", note="Replace · CMF mood / texture")
    # swatch rail left
    sy = Inches(3.25)
    sw = Inches(1.45); sh = Inches(1.45); sgap = Inches(0.25)
    for i, (name, color) in enumerate(swatches):
        x = Emu(int(Inches(0.75) + (sw + sgap) * i))
        rect(s, x, sy, sw, sh, fill=color, line=PLATINUM, line_w=Pt(0.5))
        textbox(s, x, Emu(int(sy + sh + Inches(0.12))), sw, Inches(0.4),
                [{"text": name, "size": 9.5, "color": SLATE, "spacing": 0.5,
                  "align": PP_ALIGN.LEFT}])
    textbox(s, Inches(0.75), Inches(5.7), Inches(6.0), Inches(0.5),
            [{"text": "Swatches are placeholders — replace with approved CMF palette.",
              "size": 9.5, "color": PLATINUM}])
    page_meta(s, number)
    return s


def slide_takeaways(number, kick, title, summary, points):
    """Closing takeaways — numbered, premium, minimal."""
    s = add_slide(); set_bg(s, CLOUD)
    rect(s, Inches(0), Inches(0), SLIDE_W, Inches(2.6), fill=WHITE)
    wing_mark(s, Inches(0.75), Inches(0.85), Inches(2.0))
    kicker(s, Inches(0.75), Inches(1.05), kick)
    textbox(s, Inches(0.72), Inches(1.45), Inches(11.0), Inches(0.9),
            [{"text": title, "size": 30, "color": INK, "bold": True,
              "name": FONT_DISPLAY}])
    textbox(s, Inches(0.75), Inches(2.75), Inches(11.0), Inches(0.6),
            [{"text": summary, "size": 15, "color": GRAPHITE, "line_spacing": 1.25}])
    # three numbered takeaways
    n = len(points)
    gap = Inches(0.4)
    total = Inches(11.83)
    tw = Emu(int((total - gap * (n - 1)) / n))
    for i, (head, sub) in enumerate(points):
        x = Emu(int(Inches(0.75) + (tw + gap) * i))
        textbox(s, x, Inches(3.9), tw, Inches(0.9),
                [{"text": f"{i+1:02d}", "size": 34, "color": SILVER, "bold": True,
                  "name": FONT_DISPLAY}])
        hairline(s, x, Inches(4.75), Emu(int(tw)))
        textbox(s, x, Inches(4.95), tw, Inches(1.6),
                [{"text": head, "size": 16, "color": INK, "bold": True,
                  "space_after": 6},
                 {"text": sub, "size": 12, "color": GRAPHITE, "line_spacing": 1.3}])
    page_meta(s, number)
    return s


# ----------------------------------------------------------------------------
# Build the deck
# ----------------------------------------------------------------------------

# 1 — Title
slide_title(
    1, "Stellantis · Chrysler Design",
    "Concept Portfolio Review",
    "A premium design vision for the next generation of Chrysler — "
    "presented to executive leadership.",
    "Prepared for the Executive Steering Committee · 2026",
)

# 2 — Executive Overview (strategy)
slide_strategy(
    2, "Executive Overview", "Executive Overview",
    "A focused portfolio narrative that positions Chrysler at the premium edge "
    "of modern American design.",
    [
        ("A unified design language", "One confident, refined aesthetic across the lineup."),
        ("Three flagship statements", "Chrysler 300, Airflow, and Aircross anchor the story."),
        ("Built for what's next", "A future-ready portfolio for executive endorsement."),
    ],
)

# 3 — Chrysler 300 Strategy
slide_strategy(
    3, "Chrysler 300 · Strategy", "Chrysler 300 Strategy",
    "Reassert the 300 as Chrysler's halo sedan — commanding presence with "
    "contemporary restraint.",
    [
        ("Iconic stature, modernized", "Bold proportions reinterpreted for today."),
        ("Premium without excess", "Confidence expressed through simplicity."),
        ("A flagship that leads", "Sets the tone for the entire portfolio."),
    ],
)

# 4 — Chrysler 300 Exterior (design)
slide_design(
    4, "Chrysler 300 · Exterior", "Chrysler 300 — Exterior",
    "Sculpted surfacing and a commanding stance.", layout="single",
)

# 5 — Chrysler 300 Interior (design)
slide_design(
    5, "Chrysler 300 · Interior", "Chrysler 300 — Interior",
    "A serene, driver-focused cabin in premium materials.", layout="duo",
)

# 6 — Chrysler 300 Innovation (design)
slide_design(
    6, "Chrysler 300 · Innovation", "Chrysler 300 — Innovation",
    "Signature lighting and seamless digital surfaces.", layout="triptych",
)

# 7 — Airflow Vision (strategy)
slide_strategy(
    7, "Airflow · Vision", "Airflow Vision",
    "Airflow embodies Chrysler's electric future — aerodynamic, intelligent, "
    "and effortlessly modern.",
    [
        ("Electric-first design", "Form shaped by efficiency and calm."),
        ("Intelligent by nature", "Technology that recedes into the experience."),
        ("The shape of tomorrow", "A clear signal of Chrysler's direction."),
    ],
)

# 8 — Airflow Exterior (design)
slide_design(
    8, "Airflow · Exterior", "Airflow — Exterior",
    "Aerodynamic purity and a clean, future-forward silhouette.", layout="single",
)

# 9 — Airflow Interior (design)
slide_design(
    9, "Airflow · Interior", "Airflow — Interior",
    "A lounge-like, sustainable cabin built around the occupant.", layout="duo",
)

# 10 — Aircross Strategy
slide_strategy(
    10, "Aircross · Strategy", "Aircross Strategy",
    "Aircross extends Chrysler's design language into the most important "
    "growth segment — the premium crossover.",
    [
        ("Versatile premium", "Refinement that adapts to modern life."),
        ("Confident proportions", "SUV presence with elegant restraint."),
        ("Volume with prestige", "Reach and desirability in one statement."),
    ],
)

# 11 — Aircross Exterior (design)
slide_design(
    11, "Aircross · Exterior", "Aircross — Exterior",
    "Robust yet refined — a crossover with poise.", layout="single",
)

# 12 — Aircross Interior (design)
slide_design(
    12, "Aircross · Interior", "Aircross — Interior",
    "Spacious, flexible, and quietly luxurious.", layout="duo",
)

# 13 — Chrysler Future Portfolio (portfolio)
slide_portfolio(
    13, "Portfolio · Future", "Chrysler Future Portfolio",
    "A cohesive family — today's Pacifica alongside the concepts that define "
    "what comes next.",
    [
        ("Pacifica", "Today's premium minivan benchmark."),
        ("Airflow", "The electric flagship direction."),
        ("Aircross", "The premium crossover statement."),
        ("Next Concept", "Reserved for the future vision."),
    ],
)

# 14 — Color / Material / Finish Strategy (cmf)
slide_cmf(
    14, "Design · CMF", "Color, Material & Finish",
    "A disciplined CMF strategy that signals premium craftsmanship through "
    "tone, texture, and finish — consistent across every nameplate.",
    [
        ("Platinum", SILVER),
        ("Midnight", CHRYSLER),
        ("Graphite", GRAPHITE),
        ("Linen", RGBColor(0xE7, 0xE2, 0xD8)),
        ("Pearl", RGBColor(0xF2, 0xF3, 0xF4)),
    ],
)

# 15 — Executive Takeaways
slide_takeaways(
    15, "Closing", "Executive Takeaways",
    "Three commitments to carry forward from today's review.",
    [
        ("A clear design vision", "One premium language unifying the Chrysler portfolio."),
        ("A future-ready lineup", "Flagship concepts that lead the brand forward."),
        ("Aligned for decision", "Direction set for executive endorsement and next steps."),
    ],
)

OUT = "/home/user/CLAUDE/Stellantis_Chrysler_Executive_Portfolio.pptx"
prs.save(OUT)
print(f"Saved {OUT} with {len(prs.slides._sldIdLst)} slides.")
