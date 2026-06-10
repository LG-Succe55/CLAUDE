#!/usr/bin/env python3
"""
Revise LATEST_3.pptx to align the Design Amplified deck with the
Stellantis AI Business Hub enterprise architecture
(Business Users > AI Agents > Semantic Layer > Data Products > Platforms).

Slides touched (1-based, original order):
  9   THE OPPORTUNITY ............ add "What We Learned" insight row
  10  THE CAPABILITY STACK ....... full redesign -> Design Intelligence Stack
  +   NEW after 10 ............... Design within the Stellantis AI Ecosystem
  13  THE MOONSHOT ............... add enabling-foundations strip
  18  GLOBAL KNOWLEDGE GRAPH ..... add semantic-layer alignment callout
  19  THE PROGRAM TWIN ........... inputs -> twin -> outputs flow
  25  GLOBAL DESIGN NETWORK ...... rename + roles + knowledge-flow mesh
  +   NEW before old 26 .......... Design Impact on Enterprise Value
Everything else untouched.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.ns import qn
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import etree
import copy

SRC = '/home/user/CLAUDE/work/LATEST_3.pptx'
DST = '/home/user/CLAUDE/work/LATEST_4.pptx'

# ----- design tokens (extracted from the deck) -----
LIME      = RGBColor(0xC8, 0xFF, 0x1F)
INK       = RGBColor(0x0E, 0x0E, 0x0E)   # dark text on lime
PAPER     = RGBColor(0xF4, 0xF1, 0xE8)   # primary text
GRAY      = RGBColor(0xB8, 0xB6, 0xAC)   # secondary text
MUTED     = RGBColor(0x86, 0x84, 0x7B)   # muted label
BODYGRAY  = RGBColor(0x8A, 0x8A, 0x86)   # body gray
CARD      = RGBColor(0x12, 0x12, 0x16)   # card fill
CARD_ALT  = RGBColor(0x0E, 0x0E, 0x12)   # muted card fill
CARD_LIME = RGBColor(0x19, 0x1D, 0x0A)   # lime-tinted card fill
BORDER    = RGBColor(0x2A, 0x2A, 0x30)   # card border

prs = Presentation(SRC)
SLIDES = prs.slides

# ---------------------------------------------------------------- helpers --
def _set_alpha(srgb_el, alpha_pct):
    """append <a:alpha> inside an <a:srgbClr> element. alpha_pct 0-100"""
    a = srgb_el.makeelement(qn('a:alpha'), {'val': str(int(alpha_pct * 1000))})
    srgb_el.append(a)

def fill_solid(shape, color, alpha=None):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if alpha is not None:
        srgb = shape.fill.fore_color._xFill.find(qn('a:srgbClr'))
        _set_alpha(srgb, alpha)

def line_solid(shape, color, w_emu=12700, alpha=None, dash=None):
    ln = shape.line
    ln.color.rgb = color
    ln.width = Emu(w_emu)
    if dash is not None:
        ln.dash_style = dash
    if alpha is not None:
        lnEl = ln._get_or_add_ln()
        srgb = lnEl.find(qn('a:solidFill')).find(qn('a:srgbClr'))
        _set_alpha(srgb, alpha)

def no_line(shape):
    shape.line.fill.background()

def add_card(shapes, x, y, w, h, fill=CARD, border=BORDER, border_w=12700,
             rounded=True, adj=0.075):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    sp = shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        try:
            sp.adjustments[0] = adj
        except Exception:
            pass
    if fill is None:
        sp.fill.background()
    else:
        fill_solid(sp, fill)
    if border is None:
        no_line(sp)
    else:
        line_solid(sp, border, border_w)
    sp.shadow.inherit = False
    return sp

def add_text(shapes, x, y, w, h, runs, size=12, bold=False, italic=False,
             color=PAPER, font='Arial', align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             spc=None, line_pts=None, wrap=True):
    """runs: str, or list of (text, color, bold[, italic]) tuples,
    or list of paragraph dicts {runs:[...], size, space_after,...}"""
    tb = shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    tf.auto_size = None

    if isinstance(runs, str):
        paragraphs = [{'runs': [(runs, color, bold, italic)]}]
    elif isinstance(runs, list) and runs and isinstance(runs[0], tuple):
        paragraphs = [{'runs': [(r + (bold, italic))[:4] if len(r) == 2 else
                                (r + (italic,))[:4] if len(r) == 3 else r
                                for r in runs]}]
    else:
        paragraphs = runs

    first = True
    for pdef in paragraphs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = pdef.get('align', align)
        psize = pdef.get('size', size)
        if pdef.get('space_after') is not None:
            p.space_after = Pt(pdef['space_after'])
        lp = pdef.get('line_pts', line_pts)
        if lp:
            p.line_spacing = Pt(lp)
        for rdef in pdef['runs']:
            text, rcolor = rdef[0], rdef[1]
            rbold = rdef[2] if len(rdef) > 2 else bold
            ritalic = rdef[3] if len(rdef) > 3 else italic
            r = p.add_run()
            r.text = text
            f = r.font
            f.name = pdef.get('font', font)
            f.size = Pt(pdef.get('size', psize))
            f.bold = rbold
            f.italic = ritalic
            f.color.rgb = rcolor
            rspc = pdef.get('spc', spc)
            if rspc:
                r.font._rPr.set('spc', str(rspc))
    return tb

def add_eyebrow(shapes, text, color=LIME, x=0.90, y=0.70, w=10.0, size=11):
    return add_text(shapes, x, y, w, 0.32, text, size=size, color=color,
                    font='Courier New', spc=250)

def add_line(shapes, x1, y1, x2, y2, color=LIME, w_emu=12700, alpha=50,
             dash=None, arrow_end=False, arrow_size='med'):
    conn = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                Inches(x2), Inches(y2))
    line_solid(conn, color, w_emu, alpha, dash)
    conn.shadow.inherit = False
    if arrow_end:
        lnEl = conn.line._get_or_add_ln()
        tail = lnEl.makeelement(qn('a:tailEnd'),
                                {'type': 'triangle', 'w': arrow_size, 'len': arrow_size})
        lnEl.append(tail)
    return conn

def add_dot(shapes, cx, cy, size=0.12, color=LIME, rotated=False):
    sp = shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(cx - size / 2),
                          Inches(cy - size / 2), Inches(size), Inches(size))
    fill_solid(sp, color)
    no_line(sp)
    sp.shadow.inherit = False
    if rotated:
        sp.rotation = 45
    return sp

def delete_shape(sp):
    sp._element.getparent().remove(sp._element)

def shapes_by_id(slide):
    return {sp.shape_id: sp for sp in slide.shapes}

def move_before(slide, shape, ref_element):
    """move shape element so it renders below ref_element"""
    el = shape._element
    el.getparent().remove(el)
    ref_element.addprevious(el)

def new_dark_slide():
    """blank slide on the DEFAULT layout with the deck's dark background"""
    layout = SLIDES[8].slide_layout            # DEFAULT
    slide = SLIDES.add_slide(layout)
    img_part = SLIDES[8].part.related_part('rId3')   # media/image2.png
    rid = slide.part.relate_to(img_part, RT.IMAGE)
    bg_xml = (
        '<p:bg xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        '<p:bgPr><a:blipFill dpi="0" rotWithShape="1"><a:blip r:embed="%s"><a:lum/>'
        '</a:blip><a:srcRect/><a:stretch><a:fillRect/></a:stretch></a:blipFill>'
        '<a:effectLst/></p:bgPr></p:bg>' % rid)
    bg = etree.fromstring(bg_xml)
    cSld = slide._element.find(qn('p:cSld'))
    cSld.insert(0, bg)
    return slide

def reposition_slide(slide, index):
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst):
        if sldId.get('id') == str(slide.slide_id) or sldId.rId == slide.part.partname:
            pass
    # find by relationship id
    el = None
    for sldId in sldIdLst:
        part = prs.part.related_part(sldId.rId)
        if part is slide.part:
            el = sldId
            break
    sldIdLst.remove(el)
    sldIdLst.insert(index, el)

# ===========================================================================
# SLIDE 9 - THE OPPORTUNITY : add "What We Learned"
# ===========================================================================
s9 = SLIDES[8]
ids = shapes_by_id(s9)
# shift the two wave cards + arrow down and slim them (1.50 -> 1.30 high)
for sid, top, h in [(4, 5.70, 1.30), (8, 5.70, 1.30), (7, 5.70, 1.30),
                    (5, 5.88, 0.32), (9, 5.88, 0.32),
                    (6, 6.28, 0.70), (10, 6.28, 0.70)]:
    sp = ids[sid]
    sp.top = Inches(top)
    if sid in (4, 7, 8):
        sp.height = Inches(h)

shp9 = s9.shapes
add_eyebrow(shp9, 'WHAT WE LEARNED', color=LIME, x=0.90, y=3.72, size=10)
insights = [
    ('01', 'AI tools are only the first layer.'),
    ('02', 'Enterprise value comes from reusable knowledge.'),
    ('03', 'Design must become a producer of intelligence — not only a consumer of tools.'),
    ('04', 'The future competitive advantage is institutional knowledge, not software access.'),
]
cw, gap, cx0, cy, ch = 2.80, 0.1333, 0.90, 4.12, 1.26
for i, (num, txt) in enumerate(insights):
    x = cx0 + i * (cw + gap)
    add_card(shp9, x, cy, cw, ch)
    add_text(shp9, x + 0.18, cy + 0.14, 1.0, 0.26, num, size=9.5, color=LIME,
             font='Courier New', spc=80)
    add_text(shp9, x + 0.18, cy + 0.42, cw - 0.36, ch - 0.56, txt, size=10.5,
             bold=True, color=PAPER, line_pts=13)

# ===========================================================================
# SLIDE 10 - full redesign : DESIGN INTELLIGENCE STACK
# ===========================================================================
s10 = SLIDES[9]
ids = shapes_by_id(s10)
keep = {2, 3}
for sp in list(s10.shapes):
    if sp.shape_id not in keep:
        delete_shape(sp)
# retitle
ids[2].text_frame.paragraphs[0].runs[0].text = 'DESIGN INTELLIGENCE STACK'
hl = ids[3].text_frame.paragraphs[0]
hl.runs[0].text = 'Five layers. '
hl.runs[1].text = 'One intelligence system.'

shp10 = s10.shapes
layers = [
    # label, name, items, lime?
    ('LAYER 05', 'Design Leaders & Teams', 'Judgment · taste · decision authority', True),
    ('LAYER 04', 'Design Agents',
     'Design Review Assistant · Research Assistant · Benchmark Assistant · PDP Assistant', False),
    ('LAYER 03', 'Design Semantic Layer',
     'Design vocabulary · Relationships · Brand knowledge · Design ontology', False),
    ('LAYER 02', 'Design Data Products',
     'Customer Clinics · Design Reviews · Benchmark Intelligence · Brand DNA · Program Lessons Learned', False),
    ('LAYER 01', 'Enterprise AI Platforms',
     'GDPA · Foundry · Databricks · Snowflake · AI Garage', False),
]
SX, SW, ST, RH, RG = 2.30, 9.60, 2.18, 0.78, 0.12
for i, (lab, name, items, is_lime) in enumerate(layers):
    y = ST + i * (RH + RG)
    if is_lime:
        add_card(shp10, SX, y, SW, RH, fill=LIME, border=None)
        lab_c, name_c, item_c = INK, INK, INK
    else:
        add_card(shp10, SX, y, SW, RH)
        lab_c, name_c, item_c = MUTED, PAPER, GRAY
    add_text(shp10, SX + 0.25, y + 0.11, 1.6, 0.20, lab, size=7.5, color=lab_c,
             font='Courier New', spc=80)
    add_text(shp10, SX + 0.25, y + 0.32, 3.1, 0.36, name, size=15, bold=True,
             color=name_c)
    add_text(shp10, SX + 3.40, y + 0.14, SW - 3.40 - 0.25, RH - 0.28, items,
             size=9, color=item_c, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    if i < 4:  # junction diamond between rows
        add_dot(shp10, SX + SW / 2, y + RH + RG / 2 + (0 if i % 1 else 0), 0.10,
                rotated=True)

# value-flows-up rail
rail_x = 1.78
add_line(shp10, rail_x, ST + 4 * (RH + RG) + RH, rail_x, ST, alpha=60,
         arrow_end=True)
rail_lbl = add_text(shp10, rail_x - 1.06, 4.22, 1.90, 0.26, 'VALUE FLOWS UP',
                    size=8, color=MUTED, font='Courier New', spc=200,
                    align=PP_ALIGN.CENTER, wrap=False)
rail_lbl.rotation = -90

add_text(shp10, 0.90, 6.86, 11.60, 0.38, [
    {'runs': [('Each layer powers the one above — ', BODYGRAY, False, True),
              ('from enterprise platforms to design judgment.', LIME, True, True)],
     'size': 14, 'align': PP_ALIGN.CENTER}])

# ===========================================================================
# NEW SLIDE A - DESIGN WITHIN THE STELLANTIS AI ECOSYSTEM  (after slide 10)
# ===========================================================================
sA = new_dark_slide()
shpA = sA.shapes
add_eyebrow(shpA, 'DESIGN WITHIN THE STELLANTIS AI ECOSYSTEM')
add_text(shpA, 0.85, 1.10, 11.60, 0.60, [
    {'runs': [('Design is ', PAPER, True), ('native', LIME, True),
              (' to the enterprise architecture.', PAPER, True)], 'size': 30}])

bp_x, bp_w = 0.90, 2.50          # blueprint column
de_x, de_w = 4.05, 3.55          # design column
ex_x, ex_w = 8.45, 3.85          # examples column
row_t0, row_h, row_g = 2.55, 0.64, 0.22

add_text(shpA, bp_x, 2.18, 3.4, 0.24, 'STELLANTIS ENTERPRISE ARCHITECTURE',
         size=8.5, color=MUTED, font='Courier New', spc=80, wrap=False)
add_text(shpA, de_x, 2.18, 3.4, 0.24, 'PRODUCT DESIGN OFFICE',
         size=8.5, color=LIME, font='Courier New', spc=80, wrap=False)
add_text(shpA, ex_x, 2.18, 3.9, 0.24, 'DESIGN DATA PRODUCTS · EXAMPLES',
         size=8.5, color=MUTED, font='Courier New', spc=80, wrap=False)

blueprint = ['Business Users', 'AI Agents', 'Semantic Layer', 'Data Products',
             'Enterprise AI Platforms']
design = ['Business Users', 'Design Agents', 'Design Semantic Layer',
          'Design Data Products', 'GDPA · Enterprise AI Platforms']
examples = ['Customer Clinics', 'Design Reviews', 'Benchmark Intelligence',
            'Brand DNA', 'Program Lessons Learned']

for i in range(5):
    y = row_t0 + i * (row_h + row_g)
    cyc = y + row_h / 2
    # blueprint (muted)
    add_card(shpA, bp_x, y, bp_w, row_h, fill=CARD_ALT, border=BORDER)
    add_text(shpA, bp_x + 0.20, y, bp_w - 0.36, row_h, blueprint[i], size=11,
             bold=True, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
    # mapping arrow
    add_line(shpA, bp_x + bp_w + 0.10, cyc, de_x - 0.10, cyc, alpha=45,
             arrow_end=True, arrow_size='sm')
    # design (lime-tinted)
    add_card(shpA, de_x, y, de_w, row_h, fill=CARD_LIME, border=LIME)
    add_text(shpA, de_x + 0.22, y, de_w - 0.40, row_h, design[i], size=12,
             bold=True, color=PAPER, anchor=MSO_ANCHOR.MIDDLE)
    # vertical flow arrows inside both columns
    if i < 4:
        add_line(shpA, bp_x + bp_w / 2, y + row_h + 0.02, bp_x + bp_w / 2,
                 y + row_h + row_g - 0.02, color=BORDER, alpha=None, w_emu=11430)
        add_line(shpA, de_x + de_w / 2, y + row_h + 0.02, de_x + de_w / 2,
                 y + row_h + row_g - 0.02, alpha=60, arrow_end=True,
                 arrow_size='sm')
    # examples chips
    add_card(shpA, ex_x, y, ex_w, row_h, fill=CARD, border=BORDER)
    add_dot(shpA, ex_x + 0.26, cyc, 0.10)
    add_text(shpA, ex_x + 0.46, y, ex_w - 0.60, row_h, examples[i], size=11,
             bold=True, color=PAPER, anchor=MSO_ANCHOR.MIDDLE)

# bracket linking Design Data Products row -> examples column
br_x = ex_x - 0.14
top_y = row_t0
bot_y = row_t0 + 4 * (row_h + row_g) + row_h
dp_cy = row_t0 + 3 * (row_h + row_g) + row_h / 2
add_line(shpA, br_x, top_y, br_x, bot_y, alpha=45, w_emu=9525)
add_line(shpA, de_x + de_w + 0.08, dp_cy, br_x, dp_cy, alpha=45,
         dash=MSO_LINE_DASH_STYLE.DASH)

add_text(shpA, 0.90, 7.00, 11.60, 0.36, [
    {'runs': [('Product Design Office contributes ', BODYGRAY, False, True),
              ('reusable intelligence assets', LIME, True, True),
              (' to the Stellantis AI ecosystem.', BODYGRAY, False, True)],
     'size': 14, 'align': PP_ALIGN.CENTER}])

# ===========================================================================
# SLIDE 13 - THE MOONSHOT : enabling foundations strip
# ===========================================================================
s13 = SLIDES[12]
shp13 = s13.shapes
add_text(shp13, 0.90, 6.28, 11.53, 0.24, 'THE ENABLING FOUNDATIONS',
         size=9, color=MUTED, font='Courier New', spc=200, align=PP_ALIGN.CENTER)
chips = ['DATA PRODUCTS', 'SEMANTIC LAYER', 'KNOWLEDGE GRAPHS', 'AI AGENTS']
cw, gap, ch = 2.42, 0.24, 0.46
x0 = (13.333 - (4 * cw + 3 * gap)) / 2
for i, c in enumerate(chips):
    x = x0 + i * (cw + gap)
    add_card(shp13, x, 6.62, cw, ch)
    add_dot(shp13, x + 0.26, 6.62 + ch / 2, 0.09)
    add_text(shp13, x + 0.42, 6.62, cw - 0.5, ch, c, size=9.5, color=PAPER,
             font='Courier New', spc=80, anchor=MSO_ANCHOR.MIDDLE, wrap=False)

# ===========================================================================
# SLIDE 18 - GLOBAL KNOWLEDGE GRAPH : alignment callout
# ===========================================================================
s18 = SLIDES[17]
shp18 = s18.shapes
add_card(shp18, 0.85, 5.60, 4.60, 1.06)
bar = shp18.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.87), Inches(5.70),
                      Inches(0.05), Inches(0.86))
fill_solid(bar, LIME); no_line(bar); bar.shadow.inherit = False
add_text(shp18, 1.14, 5.76, 4.12, 0.34,
         'ALIGNED WITH STELLANTIS SEMANTIC LAYER & KNOWLEDGE GRAPH STRATEGY',
         size=8.5, color=LIME, font='Courier New', spc=80, line_pts=11.5)
add_text(shp18, 1.14, 6.18, 4.12, 0.40, [
    {'runs': [('Semantic layer · knowledge graph · enterprise ontology — ', GRAY, False),
              ('one shared language of Design.', PAPER, True)], 'size': 10,
     'line_pts': 13}])

# ===========================================================================
# SLIDE 19 - THE PROGRAM TWIN : inputs -> twin -> outputs
# ===========================================================================
s19 = SLIDES[18]
ids = shapes_by_id(s19)
keep = {2, 3, 12, 13, 14, 15, 48, 49, 50, 51, 52, 53, 54, 55, 56}
for sp in list(s19.shapes):
    if sp.shape_id not in keep:
        delete_shape(sp)
shp19 = s19.shapes
glow_el = ids[12]._element  # lines must render under the hub glow

inputs = ['Design Reviews', 'Customer Clinics', 'Benchmarking',
          'Engineering Feedback', 'Manufacturing Lessons', 'Packaging Constraints']
outputs = ['Program Intelligence', 'Decision Traceability', 'Design Assistant',
           'Lessons Learned', 'PDP Acceleration']

HUB = (6.65, 3.95)
in_x, in_w, in_h = 0.95, 2.42, 0.58
in_t0, in_step = 2.05, 0.70
out_x, out_w, out_h = 9.96, 2.42, 0.62
out_t0, out_step = 2.42, 0.78

add_text(shp19, in_x, 1.72, 3.4, 0.24, 'INPUTS · WHAT THE TWIN ABSORBS',
         size=8.5, color=MUTED, font='Courier New', spc=80, wrap=False)
add_text(shp19, out_x, 2.08, 3.0, 0.24, 'OUTPUTS · WHAT THE TWIN RETURNS',
         size=8.5, color=LIME, font='Courier New', spc=80, wrap=False)

for i, name in enumerate(inputs):
    y = in_t0 + i * in_step
    cyc = y + in_h / 2
    add_card(shp19, in_x, y, in_w, in_h)
    add_text(shp19, in_x + 0.18, y, in_w - 0.30, in_h, name, size=11, bold=True,
             color=PAPER, anchor=MSO_ANCHOR.MIDDLE)
    ln = add_line(shp19, in_x + in_w + 0.04, cyc, HUB[0] - 0.92, HUB[1],
                  alpha=50, arrow_end=True, arrow_size='sm')
    move_before(s19, ln, glow_el)

for i, name in enumerate(outputs):
    y = out_t0 + i * out_step
    cyc = y + out_h / 2
    add_card(shp19, out_x, y, out_w, out_h, fill=CARD_LIME, border=LIME)
    add_text(shp19, out_x + 0.18, y, out_w - 0.30, out_h, name, size=11,
             bold=True, color=PAPER, anchor=MSO_ANCHOR.MIDDLE)
    ln = add_line(shp19, HUB[0] + 0.92, HUB[1], out_x - 0.04, cyc,
                  alpha=50, arrow_end=True, arrow_size='sm')
    move_before(s19, ln, glow_el)

# ===========================================================================
# SLIDE 25 - GLOBAL DESIGN INTELLIGENCE NETWORK
# ===========================================================================
s25 = SLIDES[24]
ids = shapes_by_id(s25)
ids[2].text_frame.paragraphs[0].runs[0].text = 'GLOBAL DESIGN INTELLIGENCE NETWORK'

# rebuild left role list
roles = [
    ('Regional AI Champions', 'drive adoption inside every studio'),
    ('AI Ambassadors', 'carry working practice across teams'),
    ('Knowledge Stewards', 'curate what the network learns'),
    ('Technology Scouts', 'scan the external frontier'),
    ('Shared Intelligence', 'one network, one memory'),
]
tf = ids[4].text_frame
tf.clear()
first = True
for name, desc in roles:
    p1 = tf.paragraphs[0] if first else tf.add_paragraph()
    first = False
    r = p1.add_run(); r.text = name
    r.font.name = 'Arial'; r.font.size = Pt(15); r.font.bold = True
    r.font.color.rgb = PAPER
    p1.space_after = Pt(2)
    p2 = tf.add_paragraph()
    r = p2.add_run(); r.text = desc
    r.font.name = 'Arial'; r.font.size = Pt(10.5); r.font.bold = False
    r.font.color.rgb = BODYGRAY
    p2.space_after = Pt(12)

# replace the spoke lines with knowledge-flow mesh
for sid in range(5, 26):
    if sid in ids:
        delete_shape(ids[sid])
shp25 = s25.shapes
nodes = {'NA': (9.40, 2.40), 'DE': (11.74, 3.40), 'FR': (11.74, 5.40),
         'IT': (9.40, 6.40), 'BR': (7.06, 5.40), 'CN': (7.06, 3.40)}
center = (9.40, 4.40)
node_ref = ids[26]._element  # GLOBAL hub circle: lines go beneath it
# spokes (faint, solid) - shared global memory
for cx, cy in nodes.values():
    ln = add_line(shp25, center[0], center[1], cx, cy, alpha=34, w_emu=9525)
    move_before(s25, ln, node_ref)
# peer ring (dashed) - knowledge flows between regions
ring = ['NA', 'DE', 'FR', 'IT', 'BR', 'CN']
for a, b in zip(ring, ring[1:] + ring[:1]):
    x1, y1 = nodes[a]; x2, y2 = nodes[b]
    ln = add_line(shp25, x1, y1, x2, y2, alpha=45, w_emu=10160,
                  dash=MSO_LINE_DASH_STYLE.DASH)
    move_before(s25, ln, node_ref)
add_text(shp25, 5.96, 7.04, 6.44, 0.30,
         [{'runs': [('Connected by knowledge flows — not reporting lines.',
                     BODYGRAY, False, True)], 'size': 11.5,
           'align': PP_ALIGN.CENTER}])

# ===========================================================================
# NEW SLIDE B - DESIGN IMPACT ON ENTERPRISE VALUE (before "WHAT SUCCESS...")
# ===========================================================================
sB = new_dark_slide()
shpB = sB.shapes
add_eyebrow(shpB, 'DESIGN IMPACT ON ENTERPRISE VALUE')
add_text(shpB, 0.85, 1.10, 9.40, 0.60, [
    {'runs': [('Design, measured in ', PAPER, True),
              ('enterprise value.', LIME, True)], 'size': 34}])
# VCP chip
add_card(shpB, 10.55, 0.62, 1.88, 0.42)
add_text(shpB, 10.55, 0.62, 1.88, 0.42, 'ALIGNED WITH VCP', size=8.5,
         color=LIME, font='Courier New', spc=80, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, wrap=False)

pillars = [
    ('01', 'PDP Acceleration',
     ['Faster decisions', 'Reduced rework', 'Improved throughput']),
    ('02', 'Product Cost',
     ['Cost-aware design decisions', 'Benchmark intelligence', 'Design simplification']),
    ('03', 'Quality',
     ['Lessons learned', 'Design issue traceability', 'Early risk detection']),
    ('04', 'Customer Experience',
     ['Customer insight intelligence', 'Brand differentiation', 'Design desirability']),
]
cw, gap, cx0, cy, ch = 2.80, 0.1333, 0.90, 2.42, 3.42
for i, (num, title, bullets) in enumerate(pillars):
    x = cx0 + i * (cw + gap)
    add_card(shpB, x, cy, cw, ch)
    add_text(shpB, x + 0.22, cy + 0.20, 1.0, 0.26, num, size=9.5, color=LIME,
             font='Courier New', spc=80)
    add_text(shpB, x + 0.22, cy + 0.50, cw - 0.44, 0.62, title, size=15,
             bold=True, color=PAPER, line_pts=17)
    add_line(shpB, x + 0.22, cy + 1.22, x + cw - 0.22, cy + 1.22, color=BORDER,
             alpha=None, w_emu=9525)
    for j, b in enumerate(bullets):
        by = cy + 1.44 + j * 0.60
        add_dot(shpB, x + 0.28, by + 0.13, 0.10)
        add_text(shpB, x + 0.50, by, cw - 0.70, 0.56, b, size=11, bold=True,
                 color=PAPER, line_pts=13.5)

add_text(shpB, 0.90, 6.42, 11.60, 0.38, [
    {'runs': [('Product Design Office contributes directly to ', BODYGRAY, False, True),
              ('enterprise value creation', LIME, True, True),
              (' through AI-enabled intelligence.', BODYGRAY, False, True)],
     'size': 14, 'align': PP_ALIGN.CENTER}])

# ===========================================================================
# reorder: A after slide 10 (idx 10), B before old slide 26
# ===========================================================================
reposition_slide(sA, 10)
reposition_slide(sB, 26)

prs.save(DST)
print('saved', DST)
print('total slides:', len(Presentation(DST).slides.__iter__.__self__._sldIdLst))
